"""S18 ONS mid-year population estimates, mid-2025 edition: read the council rows (offline).

`mye25tablesuk.xlsx` (released 1 Oct 2026) covers all four UK nations for the same reference
date, 30 June 2025. Two tables are used, both verified 24 Oct 2026:

- **MYE2 - Persons**: `Code`, `Name`, `Geography`, `All ages` (persons, all ages).
- **MYE5**: `Code`, `Name`, `Geography`, `Area (sq km)`, `Estimated Population mid-2025` — land area
  excluding inland water (Standard Area Measurements, January 2025, for Great Britain; NISRA's
  estimate for Northern Ireland; MYE notes 18 and 19).

Council rows are those whose `Geography` is a council type (361 rows). The estimates are on the
council boundaries of April 2024 (note 10), before the 2025 Barnsley–Sheffield change.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

COUNCIL_GEOGRAPHIES = frozenset({
    "Non-metropolitan District", "Unitary Authority", "Metropolitan District", "London Borough",
    "Council Area", "Local Government District",
})
MYE2 = ("MYE2 - Persons", ("Code", "Name", "Geography", "All ages"))
MYE5 = ("MYE5", ("Code", "Name", "Geography", "Area (sq km)", "Estimated Population mid-2025"))
UK = "K02000001"


class PopulationError(RuntimeError):
    """The workbook does not have the verified structure, or its tables disagree."""


@dataclass(frozen=True)
class PopulationRow:
    code: str
    name: str
    geography: str
    population: int
    land_area_km2: float


def _table(workbook, sheet: str, header: tuple[str, ...]) -> list[tuple]:
    if sheet not in workbook.sheetnames:
        raise PopulationError(f"sheet '{sheet}' is missing")
    rows = list(workbook[sheet].iter_rows(values_only=True))
    starts = [i for i, r in enumerate(rows) if r and r[0] == "Code"]
    if len(starts) != 1:
        raise PopulationError(f"{sheet}: expected one header row starting 'Code', found {len(starts)}")
    found = tuple(rows[starts[0]][: len(header)])
    if found != header:
        raise PopulationError(f"{sheet}: header {found} differs from the verified {header}")
    return [r for r in rows[starts[0] + 1:] if r and r[0]]


def read_population(path: Path) -> tuple[dict[str, PopulationRow], int]:
    """Council rows keyed by code, and the UK total. Checks MYE2 against MYE5 and the UK total."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        mye2 = _table(workbook, *MYE2)
        mye5 = _table(workbook, *MYE5)
    finally:
        workbook.close()

    uk = next((r[3] for r in mye2 if r[0] == UK), None)
    if not isinstance(uk, int):
        raise PopulationError("MYE2: United Kingdom row is missing")
    area = {r[0]: r for r in mye5}
    out: dict[str, PopulationRow] = {}
    for code, name, geography, persons, *_ in mye2:
        if geography not in COUNCIL_GEOGRAPHIES:
            continue
        if not isinstance(persons, int) or persons <= 0:
            raise PopulationError(f"MYE2: {code} population is not a positive integer ({persons!r})")
        row5 = area.get(code)
        if row5 is None:
            raise PopulationError(f"MYE5: {code} is missing")
        if row5[4] != persons:
            raise PopulationError(f"{code}: MYE2 gives {persons} but MYE5 gives {row5[4]}")
        if not isinstance(row5[3], (int, float)) or row5[3] <= 0:
            raise PopulationError(f"MYE5: {code} area is not a positive number ({row5[3]!r})")
        if code in out:
            raise PopulationError(f"MYE2: {code} appears twice")
        out[code] = PopulationRow(code=code, name=str(name), geography=str(geography),
                                  population=persons, land_area_km2=float(row5[3]))
    total = sum(r.population for r in out.values())
    if total != uk:
        raise PopulationError(f"council rows sum to {total}, the UK row is {uk}")
    return out, uk
