"""Readers for the Stage 5 small-area inputs (S15, S16, S18, S19, S20), as cached by `fetch-reference`.

Every reader checks the structure it relies on (field names, header rows, coordinate reference
system) against what was verified in the downloaded files on 24 Oct 2026, and fails loudly on any
difference instead of guessing (docs/03-DATA-SOURCES.md):

- ONS LSOA (2021) → Ward (2025) → LAD (2025) best-fit lookup, Esri JSON: LSOA21CD, LSOA21NM,
  WD25CD, WD25NM, LAD25CD, LAD25NM.
- ONS LSOA (December 2021) boundaries BGC V5, Esri JSON in EPSG:27700: LSOA21CD, LSOA21NM.
- ONS LSOA (December 2021) population-weighted centroids V4, Esri JSON points in EPSG:27700: LSOA21CD.
- Nomis Census 2021 bulk zips, LSOA CSV: TS001 (usual residents) and TS045 (households by cars or
  vans).
- MHCLG IoD 2025 File 7 CSV: LSOA code (2021), IMD rank and decile (within England).
- NISRA Data Zone 2021 shapefile in Irish Grid (EPSG:29902): DZ2021_cd, DZ2021_nm, LGD2014_cd.
- NISRA Data Zone population-weighted centroids CSV: DZ2021_code, DZ2021_name, X, Y (Irish Grid;
  the build checks every centroid lies in its Data Zone).
- NISRA Census 2021 MS-A01, sheets "DZ" and "LGD": Geography, Geography Code, All usual residents.
- NISRA SOA 2011 shapefile in EPSG:29902: SOA_CODE, SOA_LABEL.
- NISRA NIMDM 2017 SOA results (.xls), sheet "MDM": LGD2014NAME, 2015 Default Urban/Rural,
  SOA2001 (the SOA code, e.g. 95AA01S1), SOA2001_name, MDM rank and the seven domain ranks.
- NISRA Census 2021 households by car or van availability (HH_CAR_VAN_TC5, six categories), from
  the flexible table builder (provided by John): Data Zone CSV and council .xlsx.
- Leeds HASP PPFI v2.1 [LSOA] (provided by John): the 19 columns of its data dictionary; 33,755
  English and 1,917 Welsh LSOAs, 6,976 Scottish Data Zones and 890 Northern Ireland SOAs (2011).
"""

from __future__ import annotations

import csv
import io
import json
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import shapely
import xlrd
from pyogrio.raw import read


class SmallAreaInputError(RuntimeError):
    """An input does not have the structure verified for it."""


BNG = "EPSG:27700"
IRISH_GRID = "EPSG:29902"


@dataclass(frozen=True)
class Areas:
    """Polygons of one geography, in input order."""

    codes: list[str]
    names: list[str]
    geoms: np.ndarray
    crs: str
    extra: dict[str, list]


def _need(found: list[str] | tuple[str, ...], expected: list[str] | tuple[str, ...], where: str) -> None:
    missing = [c for c in expected if c not in found]
    if missing:
        raise SmallAreaInputError(f"{where}: missing {missing}; found {list(found)}")


# ---------------------------------------------------------------------------------------------
# ONS (England and Wales)


def read_lsoa_lookup(paths: list[Path]) -> dict[str, dict[str, str]]:
    """LSOA21CD → {name, ward, wardName, lad, ladName} from the best-fit lookup pages."""
    out: dict[str, dict[str, str]] = {}
    for path in paths:
        doc = json.loads(path.read_text(encoding="utf-8"))
        _need([f["name"] for f in doc["fields"]], ["LSOA21CD", "LSOA21NM", "WD25CD", "WD25NM", "LAD25CD", "LAD25NM"],
              path.name)
        for f in doc["features"]:
            a = f["attributes"]
            out[a["LSOA21CD"]] = {"name": a["LSOA21NM"], "ward": a["WD25CD"], "wardName": a["WD25NM"],
                                  "lad": a["LAD25CD"], "ladName": a["LAD25NM"]}
    return out


def read_lsoa_boundaries(paths: list[Path]) -> Areas:
    """LSOA BGC V5 pages (Esri JSON, EPSG:27700)."""
    codes, names, geoms = [], [], []
    for path in paths:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="organizePolygons", category=RuntimeWarning)
            meta, _, wkb, fields = read(path, columns=["LSOA21CD", "LSOA21NM"])
        if meta["crs"] != BNG:
            raise SmallAreaInputError(f"{path.name}: CRS is {meta['crs']}, expected {BNG}")
        codes += list(fields[0])
        names += list(fields[1])
        geoms += list(shapely.from_wkb(wkb))
    if len(set(codes)) != len(codes):
        raise SmallAreaInputError("LSOA boundary pages repeat an LSOA")
    return Areas(codes, names, np.array(geoms, dtype=object), BNG, {})


def read_lsoa_centroids(paths: list[Path]) -> dict[str, tuple[float, float]]:
    """LSOA21CD → (easting, northing) in EPSG:27700 from the PWC V4 pages."""
    out: dict[str, tuple[float, float]] = {}
    for path in paths:
        doc = json.loads(path.read_text(encoding="utf-8"))
        _need([f["name"] for f in doc["fields"]], ["LSOA21CD"], path.name)
        wkid = doc.get("spatialReference", {}).get("latestWkid") or doc.get("spatialReference", {}).get("wkid")
        if wkid != 27700:
            raise SmallAreaInputError(f"{path.name}: spatial reference {wkid}, expected 27700")
        for f in doc["features"]:
            out[f["attributes"]["LSOA21CD"]] = (float(f["geometry"]["x"]), float(f["geometry"]["y"]))
    return out


# ---------------------------------------------------------------------------------------------
# Nomis Census 2021 (England and Wales)

TS001_TOTAL = "Residence type: Total; measures: Value"
TS045_TOTAL = "Number of cars or vans: Total: All households"
TS045_NO_CAR = "Number of cars or vans: No cars or vans in household"


def _zip_csv(zip_path: Path, member: str) -> list[dict[str, str]]:
    with zipfile.ZipFile(zip_path) as zf:
        if member not in zf.namelist():
            raise SmallAreaInputError(f"{zip_path.name}: no {member}; found {zf.namelist()}")
        with zf.open(member) as fh:
            return list(csv.DictReader(io.TextIOWrapper(fh, encoding="utf-8-sig")))


def read_census_lsoa(zip_path: Path, member: str, columns: list[str]) -> dict[str, dict[str, int]]:
    """LSOA code → {column: count} from one Nomis bulk CSV."""
    rows = _zip_csv(zip_path, member)
    if not rows:
        raise SmallAreaInputError(f"{zip_path.name}/{member}: no rows")
    _need(list(rows[0].keys()), ["date", "geography code", *columns], f"{zip_path.name}/{member}")
    out: dict[str, dict[str, int]] = {}
    for r in rows:
        if r["date"] != "2021":
            raise SmallAreaInputError(f"{member}: row dated {r['date']!r}, expected 2021")
        out[r["geography code"]] = {c: int(r[c]) for c in columns}
    return out


# ---------------------------------------------------------------------------------------------
# MHCLG IoD 2025 (England)

IOD_CODE = "LSOA code (2021)"
IOD_RANK = "Index of Multiple Deprivation (IMD) Rank (where 1 is most deprived)"
IOD_DECILE = "Index of Multiple Deprivation (IMD) Decile (where 1 is most deprived 10% of LSOAs)"
IOD_LAD = "Local Authority District code (2024)"


def read_iod_file7(path: Path) -> dict[str, dict[str, int | str]]:
    """LSOA code → {rank, decile, lad2024} for England's 33,755 LSOAs."""
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        _need(reader.fieldnames or [], [IOD_CODE, IOD_RANK, IOD_DECILE, IOD_LAD], path.name)
        out = {r[IOD_CODE]: {"rank": int(r[IOD_RANK]), "decile": int(r[IOD_DECILE]), "lad2024": r[IOD_LAD]}
               for r in reader}
    return out


# ---------------------------------------------------------------------------------------------
# NISRA (Northern Ireland)


def _shapefile(zip_path: Path, layer: str, columns: list[str]):
    meta, _, wkb, fields = read(f"/vsizip/{zip_path}/{layer}.shp", columns=columns)
    if meta["crs"] != IRISH_GRID:
        raise SmallAreaInputError(f"{zip_path.name}: CRS is {meta['crs']}, expected {IRISH_GRID}")
    _need(list(meta["fields"]), columns, zip_path.name)
    return fields, shapely.from_wkb(wkb)


def read_data_zones(zip_path: Path) -> Areas:
    fields, geoms = _shapefile(zip_path, "DZ2021", ["DZ2021_cd", "DZ2021_nm", "LGD2014_cd", "LGD2014_nm"])
    return Areas(list(fields[0]), list(fields[1]), np.asarray(geoms, dtype=object), IRISH_GRID,
                 {"lgd": list(fields[2]), "lgdName": list(fields[3])})


def read_super_output_areas(zip_path: Path) -> Areas:
    fields, geoms = _shapefile(zip_path, "SOA2011", ["SOA_CODE", "SOA_LABEL"])
    return Areas(list(fields[0]), list(fields[1]), np.asarray(geoms, dtype=object), IRISH_GRID, {})


DZ_PWC_MEMBER = "census-2021-population-weighted-centroids-data-zone.csv"


def read_dz_centroids(zip_path: Path) -> dict[str, tuple[float, float]]:
    """DZ2021 code → (X, Y) in the Irish Grid, as published."""
    rows = _zip_csv(zip_path, DZ_PWC_MEMBER)
    _need(list(rows[0].keys()), ["DZ2021_code", "DZ2021_name", "X", "Y"], DZ_PWC_MEMBER)
    return {r["DZ2021_code"]: (float(r["X"]), float(r["Y"])) for r in rows}


MS_A01_HEADER = ("Geography", "Geography Code", "All usual residents")


def read_ms_a01(path: Path, sheet: str) -> dict[str, int]:
    """Geography code → all usual residents (Census 2021) from one sheet of MS-A01 ("DZ", "LGD", …)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SmallAreaInputError(f"{path.name}: no {sheet} sheet; found {wb.sheetnames}")
    out: dict[str, int] = {}
    header_seen = False
    for row in wb[sheet].iter_rows(values_only=True):
        if not header_seen:
            # The DZ sheet says "Geography Code", the LGD sheet "Geography code".
            header_seen = tuple(str(v).lower() for v in row[:3]) == tuple(h.lower() for h in MS_A01_HEADER)
            continue
        if row[1] is None:
            continue
        out[str(row[1])] = int(row[2])
    wb.close()
    if not header_seen:
        raise SmallAreaInputError(f"{path.name}: header {MS_A01_HEADER} not found in sheet {sheet}")
    return out


NIMDM_COLUMNS = (
    "LGD2014NAME", "2015 Default Urban/Rural", "SOA2001", "SOA2001_name",
    "Multiple Deprivation Measure Rank (where 1 is most deprived)",
    "Income Domain Rank (where 1 is most deprived)",
    "Employment Domain Rank (where 1 is most deprived)",
    "Health Deprivation and Disability Domain Rank (where 1 is most deprived)",
    "Education, Skills and Training Domain Rank (where 1 is most deprived)",
    "Access to Services Domain Rank (where 1 is most deprived)",
    "Living Environment Domain Rank (where 1 is most deprived)",
    "Crime and Disorder Domain Rank (where 1 is most deprived)",
)
NIMDM_DOMAINS = ("income", "employment", "health", "education", "access", "livingEnvironment", "crime")


def _clean_header(value: object) -> str:
    return " ".join(str(value).split())


def read_nimdm2017(path: Path) -> list[dict]:
    """One row per SOA from sheet "MDM": LGD name, urban/rural (2015 default), code, name, ranks."""
    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_name("MDM")
    header = [_clean_header(v) for v in sheet.row_values(0)][: len(NIMDM_COLUMNS)]
    if tuple(header) != NIMDM_COLUMNS:
        raise SmallAreaInputError(f"{path.name}: MDM header {header}, expected {list(NIMDM_COLUMNS)}")
    rows = []
    for i in range(1, sheet.nrows):
        v = sheet.row_values(i)
        if not v[2]:
            continue
        rows.append({
            "lgdName": v[0], "urbanRural2015": v[1].strip(), "code": v[2], "name": v[3],
            "mdmRank": int(v[4]),
            "domainRanks": {d: int(v[5 + k]) for k, d in enumerate(NIMDM_DOMAINS)},
        })
    return rows


# ---------------------------------------------------------------------------------------------
# NISRA car or van availability (Census 2021, households), provided by John from the table builder

CAR_VAN_CATEGORIES = {
    "0": "No cars or vans available", "1": "1 car or van available", "2": "2 cars or vans available",
    "3": "3 cars or vans available", "4": "4 cars or vans available", "5": "5 or more cars or vans available",
}


def _car_van(rows: list[tuple], where: str) -> dict[str, dict[str, int]]:
    """Area code → {households: all six categories, noCar: category 0}; every area must have all six."""
    seen: dict[str, dict[str, int]] = {}
    for code, category, label, count in rows:
        if CAR_VAN_CATEGORIES.get(str(category)) != label:
            raise SmallAreaInputError(f"{where}: category {category!r} labelled {label!r}")
        seen.setdefault(code, {})[str(category)] = int(count)
    incomplete = [c for c, v in seen.items() if set(v) != set(CAR_VAN_CATEGORIES)]
    if incomplete:
        raise SmallAreaInputError(f"{where}: areas without all six categories: {incomplete[:5]}")
    return {c: {"households": sum(v.values()), "noCar": v["0"]} for c, v in seen.items()}


DZ_CAR_HEADER = ["Census 2021 Data Zone Code", "Census 2021 Data Zone Label", "Car or Van Availability Code",
                 "Car or Van Availability Label", "Count"]


def read_car_van_dz(path: Path) -> dict[str, dict[str, int]]:
    """Households by car or van availability (HH_CAR_VAN_TC5) for each Data Zone, from the table builder CSV."""
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        if header != DZ_CAR_HEADER:
            raise SmallAreaInputError(f"{path.name}: header {header}, expected {DZ_CAR_HEADER}")
        return _car_van([(r[0], r[2], r[3], r[4]) for r in reader], path.name)


LGD_CAR_HEADER = ("Local Government District 2014 Code", "Local Government District 2014 Label",
                  "Car or Van Availability Code", "Car or Van Availability Label", "Count")


def read_car_van_lgd(path: Path) -> dict[str, dict[str, int]]:
    """The same classification for the 11 councils (sheet "Table" of the table builder's .xlsx)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows, header_seen = [], False
    for row in wb["Table"].iter_rows(values_only=True):
        if not header_seen:
            header_seen = tuple(row[:5]) == LGD_CAR_HEADER
            continue
        if row[0]:
            rows.append((row[0], row[2], row[3], row[4]))
    wb.close()
    if not header_seen:
        raise SmallAreaInputError(f"{path.name}: header {LGD_CAR_HEADER} not found")
    return _car_van(rows, path.name)


# ---------------------------------------------------------------------------------------------
# Priority Places for Food Index v2.1 (Leeds HASP, provided by John)

PPFI_DOMAINS = (
    "supermarket_proximity", "supermarket_accessibility", "ecommerce_access", "socio_demographic",
    "nonsupermarket_proximity", "food_for_families", "fuel_poverty",
)
# The 19 columns of the HASP data dictionary (record 1.0.24), in file order.
PPFI_COLUMNS = (
    ["lsoa21cd"] + [f"domain_{d}" for d in PPFI_DOMAINS] + ["country", "country_denominator", "combined"]
    + [f"pp_dec_domain_{d}" for d in PPFI_DOMAINS] + ["pp_dec_combined"]
)
PPFI_COUNTRIES = {"E": 33_755, "W": 1_917, "S": 6_976, "NI": 890}


def _rank(value: str) -> int | None:
    return int(value) if value.strip() else None


def read_ppfi(path: Path) -> dict[str, dict]:
    """Area code → {country, rank, decile, domains: {domain: {rank, decile} | None}} for every area.

    Ranks are within each country (1 = highest priority; see docs/03-DATA-SOURCES.md, S13). A domain
    with no data for an area (three domains in Northern Ireland) is None, never zero.
    """
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        if reader.fieldnames != PPFI_COLUMNS:
            raise SmallAreaInputError(f"{path.name}: columns {reader.fieldnames}, expected {PPFI_COLUMNS}")
        out: dict[str, dict] = {}
        for r in reader:
            country = r["country"]
            if PPFI_COUNTRIES.get(country) != int(r["country_denominator"]):
                raise SmallAreaInputError(f"{path.name}: {r['lsoa21cd']} country {country!r} / {r['country_denominator']}")
            domains = {}
            for d in PPFI_DOMAINS:
                rank, decile = _rank(r[f"domain_{d}"]), _rank(r[f"pp_dec_domain_{d}"])
                if (rank is None) != (decile is None):
                    raise SmallAreaInputError(f"{path.name}: {r['lsoa21cd']} {d}: rank and decile disagree on missing")
                domains[d] = None if rank is None else {"rank": rank, "decile": decile}
            out[r["lsoa21cd"]] = {"country": country, "rank": int(r["combined"]),
                                  "decile": int(r["pp_dec_combined"]), "domains": domains}
    counts = {c: sum(v["country"] == c for v in out.values()) for c in PPFI_COUNTRIES}
    if counts != PPFI_COUNTRIES:
        raise SmallAreaInputError(f"{path.name}: areas per country {counts}, expected {PPFI_COUNTRIES}")
    for c, n in PPFI_COUNTRIES.items():
        if sorted(v["rank"] for v in out.values() if v["country"] == c) != list(range(1, n + 1)):
            raise SmallAreaInputError(f"{path.name}: combined ranks in {c} are not 1…{n}")
    return out
